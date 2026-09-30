"""
Reportes consolidados (bloque 2): ventas por hora y día de la semana, categorías, matriz plato
por sucursal, compras, merma, cierre de mes, exportación a Excel y resumen semanal por push.
"""
import io
import json
from datetime import date, datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from models.audit import AuditEvent
from models.inventory_item import InventoryItem
from models.invu_sales import InvuSale, InvuSaleLine, InvuSyncDay
from models.shipment import Shipment, ShipmentItem
from models.supplier import Supplier
from models.waste import WasteItem, WasteRecord
from services import weekly_digest
from tests.conftest import auth_headers_for

# Miércoles 16 de septiembre de 2026 (día de negocio); las horas se guardan en UTC (Panamá = UTC-5).
DIA = date(2026, 9, 16)


def _venta(db, branch, invu_id, hora_local, total, dia=DIA, lineas=(), credit=False):
    cerrada = datetime(dia.year, dia.month, dia.day, 0, 15) + timedelta(hours=hora_local + 5)   # UTC (Panamá = UTC-5)
    sale = InvuSale(branch_id=branch.id, invu_order_id=invu_id, business_date=dia, opened_at=cerrada, closed_at=cerrada,
                    status="Cerrada", is_credit_note=credit, order_type="Orden Normal", subtotal=total, total=total)
    db.add(sale); db.flush()
    for i, (code, name, cat, qty, monto) in enumerate(lineas):
        db.add(InvuSaleLine(sale_id=sale.id, branch_id=branch.id, business_date=dia, invu_line_id=invu_id * 100 + i,
                            code=code, name=name, category=cat, quantity=qty, total=monto, counted=True))
    return sale


@pytest.fixture
def ventas(db_session, clayton_branch, obarrio_branch):
    _venta(db_session, clayton_branch, 1, 12, Decimal("20.00"), lineas=[("B1", "Bowl Pollo", "Bowls", 2, Decimal("20.00"))])
    _venta(db_session, clayton_branch, 2, 12, Decimal("8.00"), lineas=[("S1", "Smoothie Mango", "Smoothies", 1, Decimal("8.00"))])
    _venta(db_session, clayton_branch, 3, 19, Decimal("15.00"), dia=DIA + timedelta(days=1), lineas=[("B1", "Bowl Pollo", "Bowls", 1, Decimal("15.00"))])
    _venta(db_session, obarrio_branch, 4, 9, Decimal("10.00"), lineas=[("B1", "Bowl Pollo", "Bowls", 1, Decimal("10.00"))])
    _venta(db_session, obarrio_branch, 5, 9, Decimal("4.00"), credit=True)   # nota de crédito: no cuenta
    for b, d, tot in ((clayton_branch, DIA, "28.00"), (clayton_branch, DIA + timedelta(days=1), "15.00"), (obarrio_branch, DIA, "10.00")):
        db_session.add(InvuSyncDay(branch_id=b.id, business_date=d, orders_count=2, net_total=Decimal(tot), invu_total=Decimal(tot), matches=True))
    db_session.commit()


@pytest.fixture
def compras_y_merma(db_session, clayton_branch, obarrio_branch, supervisor_user):
    tomate = InventoryItem(name="Tomate", unit="kg", category="Vegetales")
    pollo = InventoryItem(name="Pollo", unit="kg", category="Proteínas")
    prov = Supplier(name="PriceSmart")
    db_session.add_all([tomate, pollo, prov]); db_session.commit()
    recibido = datetime(2026, 9, 16, 15, 0)   # UTC = 10:00 Panamá
    sh = Shipment(branch_id=clayton_branch.id, received_by_user_id=supervisor_user.id, received_at=recibido, supplier_id=prov.id, has_issues=True)
    sh.items.append(ShipmentItem(inventory_item_id=tomate.id, quantity=Decimal("10"), unit_cost=Decimal("2.50")))
    sh.items.append(ShipmentItem(inventory_item_id=pollo.id, quantity=Decimal("5"), unit_cost=None))
    db_session.add(sh)
    rec = WasteRecord(branch_id=obarrio_branch.id, recorded_by_user_id=supervisor_user.id, occurred_at=recibido, reason="vencido")
    rec.items.append(WasteItem(inventory_item_id=tomate.id, quantity=Decimal("2"), unit_cost=Decimal("2.50")))
    db_session.add(rec)
    db_session.commit()
    return {"tomate": tomate, "pollo": pollo}


def _h(user):
    return auth_headers_for(user)


def test_ventas_por_hora_y_dia_de_la_semana(client, ventas, admin_user):
    r = client.get(f"/api/reports/sales/time?date_from={DIA}&date_to={DIA + timedelta(days=1)}", headers=_h(admin_user))
    assert r.status_code == 200, r.text
    d = r.json()
    horas = {h["hour"]: h for h in d["by_hour"]}
    assert (horas[12]["orders"], float(horas[12]["net"])) == (2, 28.0)
    assert (horas[9]["orders"], float(horas[9]["net"])) == (1, 10.0)   # la nota de crédito no cuenta
    assert horas[19]["orders"] == 1
    dias = {w["label"]: w for w in d["by_weekday"]}
    assert dias["Miércoles"]["orders"] == 3 and dias["Miércoles"]["days"] == 1 and float(dias["Miércoles"]["avg_net_per_day"]) == 38.0
    assert dias["Jueves"]["orders"] == 1 and dias["Lunes"]["orders"] == 0


def test_categorias_y_matriz_por_sucursal(client, ventas, admin_user, clayton_branch, obarrio_branch, supervisor_user, clayton_device):
    r = client.get(f"/api/reports/sales/categories?date_from={DIA}&date_to={DIA + timedelta(days=1)}", headers=_h(admin_user)).json()
    cats = {c["category"]: c for c in r["rows"]}
    assert float(cats["Bowls"]["revenue"]) == 45.0 and float(cats["Smoothies"]["revenue"]) == 8.0
    assert cats["Bowls"]["share_pct"] == 84.9 and float(r["total_revenue"]) == 53.0

    m = client.get(f"/api/reports/sales/dish-matrix?date_from={DIA}&date_to={DIA + timedelta(days=1)}", headers=_h(admin_user)).json()
    assert [b["code"] for b in m["branches"]] == ["CLY", "OBR"]
    bowl = m["rows"][0]
    assert bowl["code"] == "B1" and float(bowl["total_qty"]) == 4.0
    assert float(bowl["by_branch"][str(clayton_branch.id)]["qty"]) == 3.0 and float(bowl["by_branch"][str(obarrio_branch.id)]["qty"]) == 1.0

    # Un supervisor de sucursal solo ve su columna.
    m = client.get(f"/api/reports/sales/dish-matrix?date_from={DIA}&date_to={DIA}", headers=auth_headers_for(supervisor_user, clayton_device.device_id)).json()
    assert [b["code"] for b in m["branches"]] == ["CLY"]


def test_compras_por_proveedor_y_merma_por_motivo(client, compras_y_merma, admin_user):
    r = client.get(f"/api/reports/purchases?date_from={DIA}&date_to={DIA}", headers=_h(admin_user)).json()
    assert float(r["total"]) == 25.0 and r["lines"] == 2 and r["lines_without_cost"] == 1
    p = r["by_supplier"][0]
    assert (p["supplier"], p["shipments"], p["lines"], p["lines_without_cost"], p["issues"]) == ("PriceSmart", 1, 2, 1, 1)
    assert r["by_category"][0]["category"] == "Vegetales" and float(r["by_category"][0]["amount"]) == 25.0

    w = client.get(f"/api/reports/waste?date_from={DIA}&date_to={DIA}", headers=_h(admin_user)).json()
    assert float(w["total"]) == 5.0
    assert w["by_reason"][0]["reason"] == "vencido" and w["by_reason"][0]["label"] == "Vencido"
    assert w["by_category"][0]["category"] == "Vegetales"


def test_cierre_de_mes(client, ventas, compras_y_merma, admin_user, clayton_branch):
    r = client.get("/api/reports/month-close?month=2026-09", headers=_h(admin_user))
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["month"] == "2026-09" and d["current"]["date_from"] == "2026-09-01"
    por = {b["branch_code"]: b for b in d["current"]["branches"]}
    cly = por["CLY"]
    assert float(cly["sales_net"]) == 43.0 and cly["orders"] == 3 and float(cly["purchases"]) == 25.0
    assert cly["purchases_pct_sales"] == 58.1 and cly["purchase_lines_without_cost"] == 1
    assert float(por["OBR"]["waste_cost"]) == 5.0 and por["OBR"]["waste_pct_sales"] == 50.0
    assert d["current"]["branches"][0]["branch_code"] == "CLY"   # la que más vendió primero
    assert float(d["current"]["total"]["sales_net"]) == 53.0 and d["previous"]["date_from"] == "2026-08-01"
    assert client.get("/api/reports/month-close?month=2026-13", headers=_h(admin_user)).status_code == 422


def test_exportaciones_a_excel(client, ventas, compras_y_merma, admin_user):
    def hoja(kind):
        r = client.get(f"/api/reports/export/{kind}.xlsx?date_from={DIA}&date_to={DIA + timedelta(days=1)}", headers=_h(admin_user))
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert f"farmhouse-{kind}-" in r.headers["content-disposition"]
        return load_workbook(io.BytesIO(r.content))

    wb = hoja("ventas")
    assert wb.sheetnames == ["Ventas por día", "Platos"]
    filas = list(wb["Ventas por día"].iter_rows(values_only=True))
    assert filas[0][0] == "Fecha" and len(filas) == 4
    assert any(f[1] == "Clayton" and f[3] == 28.0 for f in filas[1:])
    platos = list(wb["Platos"].iter_rows(values_only=True))
    assert platos[1][1] == "Bowl Pollo" and platos[1][3] == 4.0

    wb = hoja("compras")
    filas = list(wb["Compras"].iter_rows(values_only=True))
    assert len(filas) == 3 and filas[1][2] == "PriceSmart"
    assert {f[4] for f in filas[1:]} == {"Tomate", "Pollo"}

    wb = hoja("merma")
    filas = list(wb["Merma"].iter_rows(values_only=True))
    assert len(filas) == 2 and filas[1][2] == "Vencido" and filas[1][8] == 5.0

    assert hoja("conteos").sheetnames == ["Conteos"]
    assert hoja("auditoria").sheetnames == ["Auditoría"]
    assert client.get("/api/reports/export/otra.xlsx", headers=_h(admin_user)).status_code == 404


def test_solo_gerencia_ve_los_reportes(client, clayton_agent, clayton_device):
    assert client.get("/api/reports/sales/time", headers=auth_headers_for(clayton_agent, clayton_device.device_id)).status_code == 403


def test_resumen_semanal_solo_los_lunes_y_una_vez(db_session, ventas, admin_user, supervisor_user, monkeypatch):
    enviados = []
    monkeypatch.setattr("services.weekly_digest.notify_users", lambda db, ids, title, body, url, **kw: enviados.append({"ids": list(ids), "title": title, "body": body, "url": url}) or len(ids))
    lunes_21 = datetime(2026, 9, 21, 9, 0, tzinfo=weekly_digest.PANAMA_TZ)   # semana previa: 14 al 20 (incluye las ventas)

    assert weekly_digest.send_if_due(now=datetime(2026, 9, 22, 9, 0, tzinfo=weekly_digest.PANAMA_TZ), db=db_session) is False   # martes
    assert weekly_digest.send_if_due(now=datetime(2026, 9, 21, 7, 0, tzinfo=weekly_digest.PANAMA_TZ), db=db_session) is False   # muy temprano
    assert weekly_digest.send_if_due(now=lunes_21, db=db_session) is True
    assert len(enviados) == 1
    aviso = enviados[0]
    assert set(aviso["ids"]) >= {admin_user.id, supervisor_user.id}
    assert aviso["title"] == "Resumen semanal · 14/09 al 20/09"
    assert "Ventas $53" in aviso["body"] and "mejor: Clayton" in aviso["body"]
    assert aviso["url"] == "/link?date_from=2026-09-14&date_to=2026-09-20"
    ev = db_session.query(AuditEvent).filter(AuditEvent.action == "digest.weekly").one()
    assert json.loads(ev.metadata_json)["week"] == "2026-09-14"

    # La misma semana no se manda dos veces.
    assert weekly_digest.send_if_due(now=lunes_21 + timedelta(hours=3), db=db_session) is False
    assert len(enviados) == 1


def test_administracion_de_inventario(client, db_session, compras_y_merma, admin_user, clayton_branch, obarrio_branch):
    # Una línea llegó incompleta contra la factura: 12 facturados, 10 recibidos, $2.50 c/u.
    from models.shipment import ShipmentItem
    li = db_session.query(ShipmentItem).filter(ShipmentItem.unit_cost.isnot(None)).first()
    li.invoiced_quantity = Decimal("12"); li.line_status = "falto"
    db_session.commit()

    r = client.get(f"/api/reports/inventory/overview?date_from={DIA}&date_to={DIA}", headers=_h(admin_user))
    assert r.status_code == 200, r.text
    d = r.json()
    # ¿Cuánto hay? Clayton recibió 10 kg de tomate a $2.50 y 5 kg de pollo sin costo.
    cly = next(b for b in d["stock"]["branches"] if b["code"] == "CLY")
    assert cly["items_with_stock"] == 2 and float(cly["stock_value"]) == 25.0
    tomate = next(i for i in d["stock"]["items"] if i["name"] == "Tomate")
    assert float(tomate["by_branch"][str(clayton_branch.id)]["qty"]) == 10.0 and float(tomate["total_value"]) == 25.0
    assert str(obarrio_branch.id) not in tomate["by_branch"] or float(tomate["by_branch"][str(obarrio_branch.id)]["qty"]) <= 0
    # ¿Cuánto se gastó? ¿Cuánto llegó?
    assert float(d["purchases"]["total"]) == 25.0 and d["purchases"]["lines_without_cost"] == 1
    assert d["arrived"]["shipments"] == 1 and d["arrived"]["top_items"][0]["name"] == "Tomate"
    # ¿Llegó todo? No: faltaron 2 kg de tomate, reclamo $5.00.
    disc = d["discrepancies"]
    assert disc["with_issues"] == 1 and float(disc["claim_total"]) == 5.0
    assert disc["rows"][0]["supplier"] == "PriceSmart" and disc["rows"][0]["statuses"] == {"falto": 1}
    assert disc["by_supplier"][0]["with_issues"] == 1
