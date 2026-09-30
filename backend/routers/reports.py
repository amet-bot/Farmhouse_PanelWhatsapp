"""
Reportes consolidados para gerencia (bloque 2 del plan de pulido): ventas por hora, día de la
semana y categoría; matriz plato por sucursal; compras por proveedor y categoría; merma por
categoría y motivo; cierre de mes; y exportación a Excel de ventas, compras, merma, conteos y
auditoría. Todo sale ya sumado del servidor (mismas cifras que /link/sales/* e Inventario).

Acceso: admin y supervisor (reports.view); un supervisor de sucursal solo ve la suya, igual
que en /link. Las fechas son días de negocio en hora de Panamá.
"""
import io
import json
import logging
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from database import get_db
from models.audit import AuditEvent
from models.branch import Branch
from models.inventory_item import InventoryItem
from models.invu_sales import InvuSale, InvuSaleLine, InvuSyncDay
from models.ops import Incident
from models.shipment import Shipment, ShipmentItem
from models.stock_count import StockCount, StockCountItem
from models.supplier import Supplier
from models.user import User
from models.waste import WasteItem, WasteRecord
from routers.link import _gerencia, _rango, _sucursal_visible, daily_sales, item_sales
from routers.inventory import WASTE_REASON_LABELS
from services.branch_hours import PANAMA_TZ
from services.invu_sales_sync import hoy_panama

logger = logging.getLogger("farmhouse.reports")

router = APIRouter(prefix="/reports", tags=["Reportes"])

WEEKDAYS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _q(value) -> Decimal:
    return Decimal(str(value or 0)).quantize(Decimal("0.01"))


def _local(ts: Optional[datetime]) -> Optional[datetime]:
    """Las columnas DATETIME guardan UTC sin zona; para la hora del día se pasa a Panamá."""
    if ts is None:
        return None
    return ts.replace(tzinfo=timezone.utc).astimezone(PANAMA_TZ)


def _utc_bounds(desde: date, hasta: date) -> tuple:
    """Inicio y fin (exclusivo) en UTC naive de un rango de días de Panamá, para columnas como
    Shipment.received_at o WasteRecord.occurred_at que guardan el instante y no el día."""
    start = datetime.combine(desde, datetime.min.time(), tzinfo=PANAMA_TZ).astimezone(timezone.utc).replace(tzinfo=None)
    end = datetime.combine(hasta + timedelta(days=1), datetime.min.time(), tzinfo=PANAMA_TZ).astimezone(timezone.utc).replace(tzinfo=None)
    return start, end


def _pct(part: Decimal, whole: Decimal) -> Optional[float]:
    if not whole:
        return None
    return round(float(part) / float(whole) * 100, 1)


# ==========================================================================
# Ventas: hora, día de la semana, categoría, plato por sucursal
# ==========================================================================
@router.get("/sales/time")
def sales_by_time(
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    branch_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(_gerencia),
):
    """Órdenes y venta por hora del día y por día de la semana (hora de Panamá). Sirve para
    ver a qué hora se vende y qué días flojean; `avg_net_per_day` divide entre los días de ese
    nombre que hubo en el rango (un rango de 30 días tiene 4 o 5 lunes)."""
    desde, hasta = _rango(date_from, date_to, por_defecto=30)
    visible = _sucursal_visible(current_user, branch_id)
    q = db.query(InvuSale.closed_at, InvuSale.opened_at, InvuSale.business_date, InvuSale.total).filter(
        InvuSale.business_date >= desde, InvuSale.business_date <= hasta,
        InvuSale.is_credit_note == False,  # noqa: E712
    )
    if visible is not None:
        q = q.filter(InvuSale.branch_id == visible)

    by_hour = [{"hour": h, "orders": 0, "net": Decimal("0")} for h in range(24)]
    by_weekday = [{"weekday": i, "label": WEEKDAYS[i], "orders": 0, "net": Decimal("0"), "_days": set()} for i in range(7)]
    for closed_at, opened_at, business_date, total in q.all():
        monto = _q(total)
        ts = _local(closed_at or opened_at)
        if ts is not None:
            by_hour[ts.hour]["orders"] += 1
            by_hour[ts.hour]["net"] += monto
        wd = by_weekday[business_date.weekday()]
        wd["orders"] += 1
        wd["net"] += monto
        wd["_days"].add(business_date)
    for wd in by_weekday:
        dias = len(wd.pop("_days"))
        wd["days"] = dias
        wd["avg_net_per_day"] = _q(wd["net"] / dias) if dias else Decimal("0")
    return {"date_from": desde, "date_to": hasta, "by_hour": by_hour, "by_weekday": by_weekday}


@router.get("/sales/categories")
def sales_by_category(
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    branch_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(_gerencia),
):
    """Cantidad y venta por categoría del menú de Invu (Bowls, Smoothies, Toasties...)."""
    desde, hasta = _rango(date_from, date_to, por_defecto=30)
    visible = _sucursal_visible(current_user, branch_id)
    q = db.query(
        InvuSaleLine.category, func.coalesce(func.sum(InvuSaleLine.quantity), 0), func.coalesce(func.sum(InvuSaleLine.total), 0),
    ).filter(
        InvuSaleLine.business_date >= desde, InvuSaleLine.business_date <= hasta,
        InvuSaleLine.counted == True,  # noqa: E712
    )
    if visible is not None:
        q = q.filter(InvuSaleLine.branch_id == visible)
    rows = [{"category": c or "Sin categoría", "quantity": _q(qty), "revenue": _q(rev)} for c, qty, rev in q.group_by(InvuSaleLine.category).all()]
    total = sum((r["revenue"] for r in rows), Decimal("0"))
    for r in rows:
        r["share_pct"] = _pct(r["revenue"], total)
    rows.sort(key=lambda r: -r["revenue"])
    return {"date_from": desde, "date_to": hasta, "total_revenue": total, "rows": rows}


@router.get("/sales/dish-matrix")
def dish_matrix(
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    limit: int = Query(25, ge=1, le=200),
    db: Session = Depends(get_db),
    current_user: User = Depends(_gerencia),
):
    """Los platos más vendidos con una columna por sucursal: ve de un vistazo qué se vende en
    una y no en otra. Un supervisor local ve una sola columna."""
    desde, hasta = _rango(date_from, date_to, por_defecto=30)
    visible = _sucursal_visible(current_user, None)
    q = db.query(
        InvuSaleLine.code, InvuSaleLine.name, InvuSaleLine.category, InvuSaleLine.branch_id,
        func.coalesce(func.sum(InvuSaleLine.quantity), 0), func.coalesce(func.sum(InvuSaleLine.total), 0),
    ).filter(
        InvuSaleLine.business_date >= desde, InvuSaleLine.business_date <= hasta,
        InvuSaleLine.counted == True,  # noqa: E712
    )
    if visible is not None:
        q = q.filter(InvuSaleLine.branch_id == visible)
    dishes: dict = {}
    branch_ids = set()
    for code, name, category, b_id, qty, rev in q.group_by(InvuSaleLine.code, InvuSaleLine.name, InvuSaleLine.category, InvuSaleLine.branch_id).all():
        key = code or name
        d = dishes.setdefault(key, {"code": code, "name": name, "category": category, "total_qty": Decimal("0"), "total_revenue": Decimal("0"), "by_branch": {}})
        d["total_qty"] += _q(qty)
        d["total_revenue"] += _q(rev)
        cell = d["by_branch"].setdefault(b_id, {"qty": Decimal("0"), "revenue": Decimal("0")})
        cell["qty"] += _q(qty)
        cell["revenue"] += _q(rev)
        branch_ids.add(b_id)
    rows = sorted(dishes.values(), key=lambda d: -d["total_qty"])[:limit]
    for r in rows:
        r["by_branch"] = {str(k): v for k, v in r["by_branch"].items()}
    branches = db.query(Branch).filter(Branch.id.in_(branch_ids)).order_by(Branch.name).all() if branch_ids else []
    return {
        "date_from": desde, "date_to": hasta,
        "branches": [{"id": b.id, "code": b.code, "name": b.name} for b in branches],
        "rows": rows,
    }


# ==========================================================================
# Compras y merma
# ==========================================================================
def _purchase_lines(db: Session, desde: date, hasta: date, visible: Optional[int]):
    start, end = _utc_bounds(desde, hasta)
    q = db.query(ShipmentItem).join(Shipment, Shipment.id == ShipmentItem.shipment_id).options(
        joinedload(ShipmentItem.inventory_item),
        joinedload(ShipmentItem.shipment).joinedload(Shipment.supplier),
        joinedload(ShipmentItem.shipment).joinedload(Shipment.branch),
        joinedload(ShipmentItem.shipment).joinedload(Shipment.received_by_user),
    ).filter(Shipment.received_at >= start, Shipment.received_at < end)
    if visible is not None:
        q = q.filter(Shipment.branch_id == visible)
    return q.order_by(Shipment.received_at.desc()).all()


@router.get("/purchases")
def purchases(
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    branch_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(_gerencia),
):
    """Compras (cargamentos recibidos) por proveedor y por categoría de insumo. El monto suma
    cantidad × costo unitario de las líneas que traen costo; las que no lo traen se cuentan
    aparte para que se note cuánto falta por costear."""
    desde, hasta = _rango(date_from, date_to, por_defecto=30)
    visible = _sucursal_visible(current_user, branch_id)
    lines = _purchase_lines(db, desde, hasta, visible)
    por_proveedor: dict = {}
    por_categoria: dict = {}
    total = Decimal("0")
    sin_costo = 0
    for li in lines:
        sh = li.shipment
        nombre = sh.supplier.name if sh.supplier else "Sin proveedor"
        p = por_proveedor.setdefault(nombre, {"supplier": nombre, "supplier_id": sh.supplier_id, "shipments": set(), "lines": 0, "lines_without_cost": 0, "amount": Decimal("0"), "issues": set()})
        cat = (li.inventory_item.category if li.inventory_item else None) or "Sin categoría"
        c = por_categoria.setdefault(cat, {"category": cat, "lines": 0, "amount": Decimal("0")})
        p["shipments"].add(sh.id)
        p["lines"] += 1
        c["lines"] += 1
        if sh.has_issues:
            p["issues"].add(sh.id)
        if li.unit_cost is None:
            p["lines_without_cost"] += 1
            sin_costo += 1
            continue
        monto = _q(Decimal(li.quantity) * Decimal(li.unit_cost))
        p["amount"] += monto
        c["amount"] += monto
        total += monto
    proveedores = sorted(por_proveedor.values(), key=lambda p: -p["amount"])
    for p in proveedores:
        p["shipments"] = len(p["shipments"])
        p["issues"] = len(p["issues"])
        p["share_pct"] = _pct(p["amount"], total)
    categorias = sorted(por_categoria.values(), key=lambda c: -c["amount"])
    for c in categorias:
        c["share_pct"] = _pct(c["amount"], total)
    return {"date_from": desde, "date_to": hasta, "total": total, "lines": len(lines), "lines_without_cost": sin_costo, "by_supplier": proveedores, "by_category": categorias}


def _waste_lines(db: Session, desde: date, hasta: date, visible: Optional[int]):
    start, end = _utc_bounds(desde, hasta)
    q = db.query(WasteItem).join(WasteRecord, WasteRecord.id == WasteItem.waste_record_id).options(
        joinedload(WasteItem.inventory_item),
        joinedload(WasteItem.waste_record).joinedload(WasteRecord.branch),
        joinedload(WasteItem.waste_record).joinedload(WasteRecord.recorded_by_user),
    ).filter(WasteRecord.occurred_at >= start, WasteRecord.occurred_at < end)
    if visible is not None:
        q = q.filter(WasteRecord.branch_id == visible)
    return q.order_by(WasteRecord.occurred_at.desc()).all()


def _waste_cost(li: WasteItem) -> Optional[Decimal]:
    if li.unit_cost is None:
        return None
    return _q(Decimal(li.quantity) * Decimal(li.unit_cost))


@router.get("/waste")
def waste(
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    branch_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(_gerencia),
):
    """Merma en plata por categoría de insumo y por motivo (el detalle por insumo y por día
    sigue en Inventario -> Merma)."""
    desde, hasta = _rango(date_from, date_to, por_defecto=30)
    visible = _sucursal_visible(current_user, branch_id)
    lines = _waste_lines(db, desde, hasta, visible)
    por_categoria: dict = {}
    por_motivo: dict = {}
    total = Decimal("0")
    sin_costo = 0
    for li in lines:
        costo = _waste_cost(li)
        cat = (li.inventory_item.category if li.inventory_item else None) or "Sin categoría"
        motivo = li.waste_record.reason
        c = por_categoria.setdefault(cat, {"category": cat, "lines": 0, "cost": Decimal("0")})
        m = por_motivo.setdefault(motivo, {"reason": motivo, "label": WASTE_REASON_LABELS.get(motivo, motivo), "lines": 0, "cost": Decimal("0")})
        c["lines"] += 1
        m["lines"] += 1
        if costo is None:
            sin_costo += 1
            continue
        c["cost"] += costo
        m["cost"] += costo
        total += costo
    categorias = sorted(por_categoria.values(), key=lambda c: -c["cost"])
    motivos = sorted(por_motivo.values(), key=lambda m: -m["cost"])
    for r in categorias + motivos:
        r["share_pct"] = _pct(r["cost"], total)
    return {"date_from": desde, "date_to": hasta, "total": total, "lines": len(lines), "lines_without_cost": sin_costo, "by_category": categorias, "by_reason": motivos}


# ==========================================================================
# Cierre de mes
# ==========================================================================
def _month_bounds(month: str) -> tuple:
    try:
        year, mon = (int(p) for p in month.split("-"))
        desde = date(year, mon, 1)
    except (ValueError, TypeError):
        raise HTTPException(status_code=422, detail="El mes va como AAAA-MM, por ejemplo 2026-09.")
    hasta = (date(year + (mon == 12), (mon % 12) + 1, 1) - timedelta(days=1))
    return desde, min(hasta, hoy_panama())


def month_figures(db: Session, desde: date, hasta: date, visible: Optional[int]) -> dict:
    """Las cifras del cierre por sucursal (y el total) para un rango de días. Lo usa el cierre
    de mes y el resumen semanal por push (services/weekly_digest.py)."""
    branches = {b.id: b for b in db.query(Branch).filter(Branch.active == True).all()}  # noqa: E712
    ids = [visible] if visible is not None else [i for i, b in branches.items() if b.code != "CAT"]
    filas = {i: {
        "branch_id": i, "branch_code": branches[i].code if i in branches else "", "branch_name": branches[i].name if i in branches else f"#{i}",
        "sales_net": Decimal("0"), "orders": 0, "items_sold": Decimal("0"), "days_synced": 0,
        "purchases": Decimal("0"), "purchase_lines_without_cost": 0, "waste_cost": Decimal("0"), "waste_lines": 0,
        "counts": 0, "count_missing_cost": Decimal("0"), "incidents": 0,
    } for i in ids}

    for dia in db.query(InvuSyncDay).filter(InvuSyncDay.business_date >= desde, InvuSyncDay.business_date <= hasta, InvuSyncDay.net_total.isnot(None)).all():
        f = filas.get(dia.branch_id)
        if f is None:
            continue
        f["sales_net"] += _q(dia.invu_total if dia.invu_total is not None else dia.net_total)
        f["days_synced"] += 1
    for b_id, n in db.query(InvuSale.branch_id, func.count(InvuSale.id)).filter(
        InvuSale.business_date >= desde, InvuSale.business_date <= hasta, InvuSale.is_credit_note == False,  # noqa: E712
    ).group_by(InvuSale.branch_id).all():
        if b_id in filas:
            filas[b_id]["orders"] = int(n)
    for b_id, qty in db.query(InvuSaleLine.branch_id, func.coalesce(func.sum(InvuSaleLine.quantity), 0)).filter(
        InvuSaleLine.business_date >= desde, InvuSaleLine.business_date <= hasta, InvuSaleLine.counted == True,  # noqa: E712
    ).group_by(InvuSaleLine.branch_id).all():
        if b_id in filas:
            filas[b_id]["items_sold"] = _q(qty)

    for li in _purchase_lines(db, desde, hasta, visible):
        f = filas.get(li.shipment.branch_id)
        if f is None:
            continue
        if li.unit_cost is None:
            f["purchase_lines_without_cost"] += 1
        else:
            f["purchases"] += _q(Decimal(li.quantity) * Decimal(li.unit_cost))
    for li in _waste_lines(db, desde, hasta, visible):
        f = filas.get(li.waste_record.branch_id)
        if f is None:
            continue
        f["waste_lines"] += 1
        costo = _waste_cost(li)
        if costo is not None:
            f["waste_cost"] += costo

    start, end = _utc_bounds(desde, hasta)
    for c in db.query(StockCount).options(joinedload(StockCount.items)).filter(StockCount.counted_at >= start, StockCount.counted_at < end).all():
        f = filas.get(c.branch_id)
        if f is None:
            continue
        f["counts"] += 1
        for it in c.items:
            if it.difference is not None and it.difference < 0 and it.unit_cost is not None:
                f["count_missing_cost"] += _q(-Decimal(it.difference) * Decimal(it.unit_cost))
    for b_id, n in db.query(Incident.branch_id, func.count(Incident.id)).filter(Incident.created_at >= start, Incident.created_at < end).group_by(Incident.branch_id).all():
        if b_id in filas:
            filas[b_id]["incidents"] = int(n)

    def ratios(f: dict) -> dict:
        f["avg_ticket"] = _q(f["sales_net"] / f["orders"]) if f["orders"] else Decimal("0")
        f["purchases_pct_sales"] = _pct(f["purchases"], f["sales_net"])
        f["waste_pct_sales"] = _pct(f["waste_cost"], f["sales_net"])
        return f

    rows = [ratios(f) for f in filas.values()]
    rows.sort(key=lambda f: -f["sales_net"])
    total = {"branch_id": 0, "branch_code": "ALL", "branch_name": "Todas"}
    for k in ("sales_net", "purchases", "waste_cost", "count_missing_cost", "items_sold"):
        total[k] = sum((f[k] for f in rows), Decimal("0"))
    for k in ("orders", "days_synced", "purchase_lines_without_cost", "waste_lines", "counts", "incidents"):
        total[k] = sum(f[k] for f in rows)
    ratios(total)
    return {"date_from": desde, "date_to": hasta, "branches": rows, "total": total}


@router.get("/month-close")
def month_close(
    month: Optional[str] = Query(None, pattern=r"^\d{4}-\d{2}$"),
    branch_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(_gerencia),
):
    """Cierre de mes por sucursal: venta, órdenes, ticket, compras y merma con su porcentaje
    sobre la venta, faltantes de conteo en plata e incidencias, comparado con el mes anterior."""
    hoy = hoy_panama()
    month = month or f"{hoy.year}-{hoy.month:02d}"
    desde, hasta = _month_bounds(month)
    prev_hasta = desde - timedelta(days=1)
    prev_desde = prev_hasta.replace(day=1)
    visible = _sucursal_visible(current_user, branch_id)
    actual = month_figures(db, desde, hasta, visible)
    anterior = month_figures(db, prev_desde, prev_hasta, visible)
    return {"month": month, "current": actual, "previous": anterior}


# ==========================================================================
# Exportación a Excel
# ==========================================================================
def _workbook(sheets: list) -> bytes:
    """sheets: [(título, encabezados, filas)] -> bytes .xlsx. Encabezado en negrita, primera
    fila fija y anchos según el contenido."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F8F6A")
        for row in rows:
            ws.append([float(v) if isinstance(v, Decimal) else (v.isoformat() if isinstance(v, (date, datetime)) else v) for v in row])
        ws.freeze_panes = "A2"
        for idx, header in enumerate(headers, start=1):
            largo = max([len(str(header))] + [len(str(r[idx - 1])) for r in rows[:200] if idx - 1 < len(r) and r[idx - 1] is not None])
            ws.column_dimensions[get_column_letter(idx)].width = min(max(10, largo + 2), 48)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _fecha_local(ts: Optional[datetime]) -> str:
    loc = _local(ts)
    return loc.strftime("%Y-%m-%d %H:%M") if loc else ""


@router.get("/export/{kind}.xlsx")
def export_excel(
    kind: str,
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    branch_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(_gerencia),
):
    """Un archivo de Excel por tema: ventas (por día y por plato), compras (línea por línea),
    merma, conteos y auditoría. Mismos filtros de fecha y sucursal que las pantallas."""
    desde, hasta = _rango(date_from, date_to, por_defecto=30)
    visible = _sucursal_visible(current_user, branch_id)

    if kind == "ventas":
        dias = daily_sales(date_from=desde, date_to=hasta, branch_id=branch_id, db=db, current_user=current_user)
        platos = item_sales(date_from=desde, date_to=hasta, branch_id=branch_id, limit=500, db=db, current_user=current_user)
        sheets = [
            ("Ventas por día", ["Fecha", "Sucursal", "Órdenes", "Venta neta (Invu)", "Venta antes de descuentos", "Descuentos", "Platos vendidos", "Cuadra con Invu"],
             [[r.business_date, r.branch_name, r.orders_count, r.net_total, r.gross_total, r.discount_total, r.items_sold, "Sí" if r.matches else "No"] for r in dias]),
            ("Platos", ["Código", "Plato", "Categoría", "Cantidad", "Venta", "Sucursales"],
             [[r.code, r.name, r.category, r.quantity, r.revenue, r.branches] for r in platos]),
        ]
    elif kind == "compras":
        rows = []
        for li in _purchase_lines(db, desde, hasta, visible):
            sh = li.shipment
            item = li.inventory_item
            rows.append([
                _fecha_local(sh.received_at), sh.branch.name if sh.branch else "", sh.supplier.name if sh.supplier else "", sh.invoice_number or "",
                item.name if item else "", item.category if item else "", li.quantity, item.unit if item else "",
                li.unit_cost, _q(Decimal(li.quantity) * Decimal(li.unit_cost)) if li.unit_cost is not None else None,
                li.invoiced_quantity, li.line_status or "ok", sh.received_by_user.name if sh.received_by_user else "",
            ])
        sheets = [("Compras", ["Recibido", "Sucursal", "Proveedor", "Factura", "Insumo", "Categoría", "Cantidad", "Unidad", "Costo unitario", "Total", "Facturado", "Estado línea", "Recibió"], rows)]
    elif kind == "merma":
        rows = []
        for li in _waste_lines(db, desde, hasta, visible):
            rec = li.waste_record
            item = li.inventory_item
            rows.append([
                _fecha_local(rec.occurred_at), rec.branch.name if rec.branch else "", WASTE_REASON_LABELS.get(rec.reason, rec.reason),
                item.name if item else "", item.category if item else "", li.quantity, item.unit if item else "",
                li.unit_cost, _waste_cost(li), rec.recorded_by_user.name if rec.recorded_by_user else "", rec.notes or "",
            ])
        sheets = [("Merma", ["Fecha", "Sucursal", "Motivo", "Insumo", "Categoría", "Cantidad", "Unidad", "Costo unitario", "Costo", "Registró", "Notas"], rows)]
    elif kind == "conteos":
        start, end = _utc_bounds(desde, hasta)
        q = db.query(StockCount).options(joinedload(StockCount.items).joinedload(StockCountItem.inventory_item), joinedload(StockCount.branch), joinedload(StockCount.counted_by_user)).filter(StockCount.counted_at >= start, StockCount.counted_at < end)
        if visible is not None:
            q = q.filter(StockCount.branch_id == visible)
        rows = []
        for c in q.order_by(StockCount.counted_at.desc()).all():
            for it in c.items:
                item = it.inventory_item
                rows.append([
                    _fecha_local(c.counted_at), c.branch.name if c.branch else "", item.name if item else "", item.unit if item else "",
                    it.expected_quantity, it.counted_quantity, it.difference, it.unit_cost,
                    _q(Decimal(it.difference) * Decimal(it.unit_cost)) if it.unit_cost is not None and it.difference is not None else None,
                    c.counted_by_user.name if c.counted_by_user else "",
                ])
        sheets = [("Conteos", ["Fecha", "Sucursal", "Insumo", "Unidad", "Esperado", "Contado", "Diferencia", "Costo unitario", "Valor diferencia", "Contó"], rows)]
    elif kind == "auditoria":
        start, end = _utc_bounds(desde, hasta)
        q = db.query(AuditEvent, User.name, Branch.name).outerjoin(User, User.id == AuditEvent.actor_user_id).outerjoin(Branch, Branch.id == AuditEvent.branch_id).filter(AuditEvent.created_at >= start, AuditEvent.created_at < end)
        if visible is not None:
            q = q.filter(AuditEvent.branch_id == visible)
        rows = [[_fecha_local(e.created_at), actor or "Sistema", branch or "", e.action, e.entity_type, e.entity_id, e.metadata_json or ""] for e, actor, branch in q.order_by(AuditEvent.created_at.desc()).limit(20000).all()]
        sheets = [("Auditoría", ["Fecha", "Usuario", "Sucursal", "Acción", "Entidad", "ID", "Detalle"], rows)]
    else:
        raise HTTPException(status_code=404, detail="Exportación no disponible. Opciones: ventas, compras, merma, conteos, auditoria.")

    contenido = _workbook(sheets)
    nombre = f"farmhouse-{kind}-{desde}-{hasta}.xlsx"
    return StreamingResponse(io.BytesIO(contenido), media_type=EXCEL_MIME, headers={"Content-Disposition": f'attachment; filename="{nombre}"'})
